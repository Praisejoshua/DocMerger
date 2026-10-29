import io
import struct
import tempfile
import zlib
from pathlib import Path

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.shared import Inches

from .services import MergeError, merge_documents


def make_png(width=4, height=4, rgb=(47, 111, 237)) -> bytes:
    raw = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))

    def chunk(tag, data):
        body = tag + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def make_docx(title: str, with_table=True, with_image=True, bullets=("one", "two")) -> bytes:
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(f"Body text of {title}.")
    for item in bullets:
        doc.add_paragraph(f"{title} {item}", style="List Bullet")
    if with_table:
        table = doc.add_table(rows=2, cols=2)
        table.style = "Table Grid"
        table.cell(0, 0).text = f"{title} cell"
    if with_image:
        doc.add_picture(io.BytesIO(make_png()), width=Inches(1))
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


class MergeServiceTests(TestCase):
    def test_merge_keeps_order_tables_and_images(self):
        sources = [(f"{t}.docx", make_docx(t)) for t in ("Alpha", "Beta", "Gamma")]
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out.docx"
            merge_documents(sources, out)
            merged = Document(str(out))

        texts = [p.text for p in merged.paragraphs]
        positions = [texts.index(t) for t in ("Alpha", "Beta", "Gamma")]
        self.assertEqual(positions, sorted(positions))
        self.assertEqual(len(merged.tables), 3)
        self.assertEqual(len(merged.inline_shapes), 3)
        for blip in merged.element.body.iter(qn("a:blip")):
            rid = blip.get(qn("r:embed"))
            self.assertEqual(merged.part.rels[rid].reltype, RT.IMAGE)

    def test_each_document_gets_its_own_section(self):
        sources = [(f"{t}.docx", make_docx(t)) for t in ("Alpha", "Beta")]
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out.docx"
            merge_documents(sources, out)
            merged = Document(str(out))
        self.assertEqual(len(merged.sections), 2)

    def test_rejects_non_docx_content(self):
        good = make_docx("Alpha")
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(MergeError):
                merge_documents([("a.docx", good), ("b.docx", b"not a word file")], Path(tmp) / "x.docx")

    def test_rejects_legacy_or_encrypted_files(self):
        good = make_docx("Alpha")
        fake_ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(MergeError) as ctx:
                merge_documents([("a.docx", good), ("b.docx", fake_ole)], Path(tmp) / "x.docx")
        self.assertIn("password", str(ctx.exception))


class MergeViewTests(TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        override = override_settings(DOCMERGE_STORAGE_DIR=Path(self.tmp.name))
        override.enable()
        self.addCleanup(override.disable)

    @staticmethod
    def upload(name, data):
        return SimpleUploadedFile(name, data, content_type="application/octet-stream")

    def test_home_page_renders(self):
        response = self.client.get(reverse("merger:home"))
        self.assertContains(response, "DocMerge")
        self.assertContains(response, "Drag and drop your Word documents here")

    def test_requires_two_files(self):
        response = self.client.post(reverse("merger:merge"), {"documents": [self.upload("a.docx", make_docx("A"))]})
        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.json()["ok"])

    def test_rejects_wrong_extension(self):
        files = [self.upload("a.docx", make_docx("A")), self.upload("notes.pdf", b"%PDF-1.4")]
        response = self.client.post(reverse("merger:merge"), {"documents": files})
        self.assertEqual(response.status_code, 400)
        self.assertIn("notes.pdf", response.json()["error"])

    def test_full_merge_and_download_flow(self):
        files = [self.upload("a.docx", make_docx("Alpha")), self.upload("b.docx", make_docx("Beta"))]
        response = self.client.post(reverse("merger:merge"), {"documents": files})
        self.assertEqual(response.status_code, 200, response.content)
        payload = response.json()
        self.assertTrue(payload["ok"])

        download = self.client.get(payload["download_url"])
        self.assertEqual(download.status_code, 200)
        content = b"".join(download.streaming_content)
        self.assertTrue(content.startswith(b"PK"))
        merged = Document(io.BytesIO(content))
        self.assertIn("Alpha", [p.text for p in merged.paragraphs])
        self.assertIn("Beta", [p.text for p in merged.paragraphs])

        self.client.post(payload["discard_url"])
        self.assertEqual(self.client.get(payload["download_url"]).status_code, 404)

    def test_other_sessions_cannot_download(self):
        files = [self.upload("a.docx", make_docx("A")), self.upload("b.docx", make_docx("B"))]
        payload = self.client.post(reverse("merger:merge"), {"documents": files}).json()
        stranger = self.client_class()
        self.assertEqual(stranger.get(payload["download_url"]).status_code, 404)


class ManyFilesTests(TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        override = override_settings(DOCMERGE_STORAGE_DIR=Path(self.tmp.name))
        override.enable()
        self.addCleanup(override.disable)

    def test_three_files_merge_in_order(self):
        names = ("Alpha", "Beta", "Gamma")
        files = [SimpleUploadedFile(f"{n}.docx", make_docx(n)) for n in names]
        payload = self.client.post(reverse("merger:merge"), {"documents": files}).json()
        self.assertTrue(payload["ok"], payload)
        content = b"".join(self.client.get(payload["download_url"]).streaming_content)
        texts = [p.text for p in Document(io.BytesIO(content)).paragraphs]
        positions = [texts.index(n) for n in names]
        self.assertEqual(positions, sorted(positions))

    def test_error_says_how_many_files_arrived(self):
        one = [SimpleUploadedFile("a.docx", make_docx("A"))]
        payload = self.client.post(reverse("merger:merge"), {"documents": one}).json()
        self.assertIn("the server received 1", payload["error"])
